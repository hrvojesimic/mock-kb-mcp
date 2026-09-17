"""MCP tools for Excel workbooks in named workspaces."""

from contextlib import closing
import os
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from fastmcp import FastMCP
from fastmcp.exceptions import ToolError
from openpyxl import load_workbook

mcp = FastMCP("Excel")


def workspace_directories() -> dict[str, Path]:
    """Look up available workspaces by ID."""
    shared_root = os.environ.get("KB_EXCEL_ROOT", "data")
    root = Path(shared_root).resolve()
    if not root.is_dir():
        raise ToolError("Workspaces are currently unavailable.")
    workspaces = {}
    for directory in sorted(root.iterdir()):
        if not directory.is_dir():
            continue
        if not directory.resolve().is_relative_to(root):
            raise ToolError("Workspace is unavailable.")
        id_path = directory / "id.txt"
        if not id_path.exists():
            try:
                with id_path.open("x", encoding="utf-8") as id_file:
                    id_file.write(str(uuid4()) + "\n")
            except FileExistsError:
                pass
        try:
            workspace_id = str(UUID(id_path.read_text(encoding="utf-8").strip()))
        except ValueError as exc:
            raise ToolError(f"Invalid workspace ID for workspace {directory.name}.") from exc
        if workspace_id in workspaces:
            raise ToolError(f"Duplicate workspace ID: {workspace_id}.")
        workspaces[workspace_id] = directory
    return workspaces


@mcp.tool(annotations={"readOnlyHint": False, "destructiveHint": False})
def list_workspaces() -> dict[str, str]:
    """List available workspaces as a mapping of workspace IDs to workspace names.

    Use a returned workspace ID when calling the Excel tools.
    """
    return {workspace_id: directory.name for workspace_id, directory in workspace_directories().items()}


def workbook_path(workspace_id: str, path: str) -> Path:
    """Locate an XLSX workbook in the selected workspace."""
    directory = workspace_directories().get(workspace_id)
    if directory is None:
        raise ToolError(f"Unknown workspace ID: {workspace_id}.")
    root = directory.resolve()
    candidate = Path(path)
    if not candidate.is_absolute():
        candidate = root / candidate
    candidate = candidate.resolve()
    if not candidate.is_relative_to(root):
        raise ToolError("Workbook path is outside the workspace.")
    if candidate.suffix.lower() != ".xlsx":
        raise ToolError("Only .xlsx workbooks are supported.")
    if not candidate.is_file():
        raise ToolError("Workbook not found in the selected workspace.")
    return candidate


@mcp.tool(annotations={"readOnlyHint": True, "destructiveHint": False})
def list_excel_sheets(workspace_id: str, path: str) -> list[str]:
    """List worksheet names in an XLSX workbook in the selected workspace.

    Use a workspace_id from list_workspaces and the workbook's path within that workspace.
    """
    with closing(load_workbook(workbook_path(workspace_id, path), read_only=True, data_only=True)) as workbook:
        return workbook.sheetnames


@mcp.tool(annotations={"readOnlyHint": True, "destructiveHint": False})
def read_excel_rows(
    workspace_id: str, path: str, sheet: str, start_row: int = 1, limit: int = 100
) -> list[dict[str, Any]]:
    """Read up to limit rows (1–1000), starting at a 1-based Excel row number.

    Returns row numbers and cell values, including headers and blank cells.
    Formulas return their cached values; formulas are not calculated.
    Use a workspace_id from list_workspaces and the workbook's path within that workspace.
    """
    if start_row < 1 or not 1 <= limit <= 1000:
        raise ValueError("start_row must be >= 1 and limit must be between 1 and 1000")
    with closing(load_workbook(workbook_path(workspace_id, path), read_only=True, data_only=True)) as workbook:
        worksheet = workbook[sheet]
        if worksheet.max_row is not None and start_row > worksheet.max_row:
            return []
        end_row = start_row + limit - 1
        if worksheet.max_row is not None:
            end_row = min(end_row, worksheet.max_row)
        return [
            {"row": number, "values": list(values)}
            for number, values in enumerate(
                worksheet.iter_rows(
                    min_row=start_row, max_row=end_row, values_only=True
                ),
                start=start_row,
            )
        ]


@mcp.tool(annotations={"readOnlyHint": True, "destructiveHint": False})
def search_excel(
    workspace_id: str, path: str, query: str, sheet: str | None = None, limit: int = 100
) -> list[dict[str, Any]]:
    """Find cells containing query (case-insensitive) in one or all worksheets.

    Returns up to limit matches (1–1000), with sheet, cell address and value.
    Searches cached formula values, without calculating formulas.
    Use a workspace_id from list_workspaces and the workbook's path within that workspace.
    """
    if not query or not 1 <= limit <= 1000:
        raise ValueError("query must not be empty and limit must be between 1 and 1000")
    matches = []
    needle = query.casefold()
    with closing(load_workbook(workbook_path(workspace_id, path), read_only=True, data_only=True)) as workbook:
        worksheets = [workbook[sheet]] if sheet is not None else workbook.worksheets
        for worksheet in worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None and needle in str(cell.value).casefold():
                        matches.append(
                            {"sheet": worksheet.title, "cell": cell.coordinate, "value": cell.value}
                        )
                        if len(matches) >= limit:
                            return matches
    return matches


if __name__ == "__main__":
    mcp.run()
